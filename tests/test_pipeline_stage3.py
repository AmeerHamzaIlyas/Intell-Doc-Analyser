"""Unit tests for Stage 3: Extraction, Cleaning, Structure Detection, Chunking, and Ingestion."""
import io
import shutil
import tempfile
from pathlib import Path
import docx
import pypdf
import pytest
from app.core.exceptions import (
    ExtractionError,
    FileCorruptedError,
    FileSizeExceededError,
    UnsupportedFileTypeError,
    ValidationError,
)
from app.services.document_service import DocumentService
from app.services.extraction.docx_extractor import DocxExtractor
from app.services.extraction.pdf_extractor import PDFExtractor
from app.services.extraction.text_extractor import TextExtractor
from app.services.processing.cleaner import TextCleaner
from app.services.processing.chunker import IntelligentChunker
from app.services.processing.structure import StructureDetector
from app.services.storage.database import Database


@pytest.fixture
def stage3_env():
    """Create isolated temporary environment for Stage 3 tests."""
    temp_dir = Path(tempfile.mkdtemp())
    db_path = temp_dir / "stage3_test.db"
    test_db = Database(db_path=db_path)
    test_service = DocumentService(db_instance=test_db)
    
    yield {
        "dir": temp_dir,
        "db": test_db,
        "service": test_service,
    }
    shutil.rmtree(temp_dir, ignore_errors=True)


def test_validation_rules(stage3_env):
    """Test file size, extension, and magic bytes validation."""
    service: DocumentService = stage3_env["service"]

    # 1. Empty file
    with pytest.raises(ValidationError):
        service.validate_file("empty.txt", b"")

    # 2. Unsupported extension
    with pytest.raises(UnsupportedFileTypeError):
        service.validate_file("script.py", b"print('hello')")

    # 3. Corrupt PDF magic bytes
    with pytest.raises(FileCorruptedError):
        service.validate_file("fake.pdf", b"This is not a pdf file")

    # 4. Valid PDF magic bytes
    ext, mime = service.validate_file("real.pdf", b"%PDF-1.4\n%real content")
    assert ext == ".pdf"
    assert mime == "application/pdf"


def test_text_cleaner():
    """Test Unicode NFKC, hyphenation repair, and whitespace normalization."""
    raw = "The deep learn-\ning architec-\nture was fast.\n\n\n\nIt processed   data."
    cleaned = TextCleaner.clean_text(raw)
    assert "learning" in cleaned
    assert "architecture" in cleaned
    assert "\n\n\n" not in cleaned
    assert "   " not in cleaned


def test_text_extractor(stage3_env):
    """Test extraction from plaintext and Markdown files."""
    temp_dir = stage3_env["dir"]
    txt_file = temp_dir / "sample.md"
    content = "# Artificial Intelligence\n\nMachine learning is a subset of AI.\n\n## Deep Learning\nNeural networks."
    txt_file.write_text(content, encoding="utf-8")

    extractor = TextExtractor()
    extracted = extractor.extract(txt_file)

    assert extracted.total_pages == 1
    assert extracted.title == "Artificial Intelligence"
    assert "Machine learning" in extracted.full_text
    assert len(extracted.pages[0].section_hints) == 2


def test_docx_extractor_with_table(stage3_env):
    """Test extraction of headings, paragraphs, and tables from DOCX."""
    temp_dir = stage3_env["dir"]
    docx_file = temp_dir / "sample.docx"

    # Create synthetic DOCX
    doc = docx.Document()
    doc.add_heading("Financial Summary 2026", level=1)
    doc.add_paragraph("This paragraph provides fiscal overview.")

    # Add a table
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Quarter"
    table.rows[0].cells[1].text = "Revenue"
    table.rows[1].cells[0].text = "Q1"
    table.rows[1].cells[1].text = "$10M"

    doc.save(str(docx_file))

    extractor = DocxExtractor()
    extracted = extractor.extract(docx_file)

    assert extracted.title == "Financial Summary 2026"
    assert "Fiscal overview" in extracted.full_text or "fiscal overview" in extracted.full_text
    # Check Markdown table conversion
    assert "| Quarter | Revenue |" in extracted.full_text
    assert "| Q1 | $10M |" in extracted.full_text


def test_pdf_extractor_multi_page(stage3_env):
    """Test extraction from a multi-page PDF."""
    temp_dir = stage3_env["dir"]
    pdf_file = temp_dir / "multipage.pdf"

    # Create a 2-page PDF using pypdf
    writer = pypdf.PdfWriter()
    page1 = writer.add_blank_page(width=300, height=300)
    page2 = writer.add_blank_page(width=300, height=300)
    
    with open(pdf_file, "wb") as f:
        writer.write(f)

    extractor = PDFExtractor()
    extracted = extractor.extract(pdf_file)
    assert extracted.total_pages == 2
    assert len(extracted.pages) == 2


def test_intelligent_chunker():
    """Test intelligent structure-aware chunker preserves provenance and respects boundaries."""
    chunker = IntelligentChunker(chunk_size=120, chunk_overlap=30)
    
    from app.services.extraction.base import ExtractedPage
    pages = [
        ExtractedPage(
            page_number=1,
            text="# Project Overview\n\nThis is the first paragraph describing document intelligence.\n\n# Architecture\n\nThe backend is built with FastAPI and SQLite.",
            tables=[],
            section_hints=["Project Overview", "Architecture"],
        )
    ]
    
    chunks = chunker.chunk_document("doc_test", pages)
    assert len(chunks) >= 2
    assert all(c.document_id == "doc_test" for c in chunks)
    assert all(c.page_number == 1 for c in chunks)
    assert all(c.section_title in ("Project Overview", "Architecture", "General") for c in chunks)
    assert all(c.token_count > 0 for c in chunks)


def test_document_service_end_to_end(stage3_env):
    """Test full DocumentService pipeline from ingestion to database storage and deduplication."""
    service: DocumentService = stage3_env["service"]
    db: Database = stage3_env["db"]

    sample_md = b"# Quarterly Report\n\nCompany metrics are growing exponentially.\n\n## Product Metrics\nActive users reached 100k."
    
    # First ingestion
    metadata, chunks = service.ingest_document("quarterly_report.md", sample_md)
    assert metadata.status == "INDEXED"
    assert metadata.word_count > 10
    assert len(chunks) >= 1
    assert db.get_document(metadata.document_id) is not None

    # Deduplication test: ingest same content again
    dup_meta, dup_chunks = service.ingest_document("quarterly_report_copy.md", sample_md)
    assert dup_meta.document_id == metadata.document_id
    assert len(dup_chunks) == len(chunks)


# ---------------------------------------------------------------------------
# Test Helpers for Synthetic Documents
# ---------------------------------------------------------------------------

def _create_synthetic_pdf(text: str = "Intelligent Document Understanding", num_pages: int = 1) -> bytes:
    """Generate valid PDF bytes with text streams and correct cross-references."""
    writer = pypdf.PdfWriter()
    for i in range(num_pages):
        page_text = f"{text} - Section {i + 1}"
        stream_content = f"BT /F1 12 Tf 50 700 Td ({page_text}) Tj ET".encode("ascii")
        length = len(stream_content)
        raw_pdf = (
            b"%PDF-1.4\n"
            b"1 0 obj <</Type /Catalog /Pages 2 0 R>> endobj\n"
            b"2 0 obj <</Type /Pages /Kids [3 0 R] /Count 1>> endobj\n"
            b"3 0 obj <</Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources <</Font <</F1 5 0 R>>>>>> endobj\n"
            b"4 0 obj <</Length " + str(length).encode() + b">> stream\n" + stream_content + b"\nendstream\nendobj\n"
            b"5 0 obj <</Type /Font /Subtype /Type1 /BaseFont /Helvetica>> endobj\n"
            b"trailer <</Size 6 /Root 1 0 R>>\nstartxref\n0\n%%EOF\n"
        )
        r = pypdf.PdfReader(io.BytesIO(raw_pdf))
        writer.add_page(r.pages[0])
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _create_encrypted_pdf(password: str = "secret123") -> bytes:
    """Generate an encrypted, password-protected PDF."""
    pdf_bytes = _create_synthetic_pdf("Confidential Financial Statement")
    reader = pypdf.PdfReader(io.BytesIO(pdf_bytes))
    writer = pypdf.PdfWriter()
    writer.add_page(reader.pages[0])
    writer.encrypt(password)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Rigorous Document Pipeline Tests (Valid, Corrupted, Empty, Scanned, Large,
# DOCX, TXT, Unsupported, Malformed, Encrypted, Unicode)
# ---------------------------------------------------------------------------

def test_valid_pdf_pipeline(stage3_env):
    """Test end-to-end ingestion and chunking of a valid PDF."""
    service: DocumentService = stage3_env["service"]
    pdf_bytes = _create_synthetic_pdf("Enterprise AI Document Extraction Pipeline", 1)
    
    meta, chunks = service.ingest_document("enterprise_report.pdf", pdf_bytes)
    assert meta.status == "INDEXED"
    assert meta.page_count == 1
    assert meta.mime_type == "application/pdf"
    assert meta.file_type == "pdf"
    assert meta.word_count > 0
    assert len(chunks) >= 1
    assert "Enterprise" in chunks[0].text
    assert "statistics" in meta.extracted_entities
    stats = meta.extracted_entities["statistics"]
    assert stats["word_count"] > 0
    assert stats["char_count"] > 0


def test_corrupted_pdf_pipeline(stage3_env):
    """Test corrupted PDF bytes safely raise FileCorruptedError and record FAILED status."""
    service: DocumentService = stage3_env["service"]
    db: Database = stage3_env["db"]
    corrupted_bytes = b"%PDF-1.4\nGarbageBinaryDataThatCausesPdfReaderToCrash\x00\xff\xfe"
    
    with pytest.raises(FileCorruptedError):
        service.ingest_document("damaged_invoice.pdf", corrupted_bytes)
    
    # Verify failed record in database
    docs = db.list_documents()
    assert len(docs) == 1
    assert docs[0].status == "FAILED"
    assert docs[0].error_message is not None


def test_empty_pdf_pipeline(stage3_env):
    """Test detection and handling of empty PDF with 0 extractable text without crashing."""
    service: DocumentService = stage3_env["service"]
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    empty_pdf_bytes = buf.getvalue()
    
    meta, chunks = service.ingest_document("blank_canvas.pdf", empty_pdf_bytes)
    assert meta.status == "INDEXED"
    assert meta.page_count == 1
    assert meta.word_count == 0
    assert meta.char_count == 0
    warnings = meta.extracted_entities.get("quality_warnings", [])
    assert any("EMPTY_DOCUMENT" in w for w in warnings)


def test_scanned_pdf_pipeline(stage3_env):
    """Test scanned document processing path when text layer is minimal."""
    service: DocumentService = stage3_env["service"]
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=300, height=300)
    buf = io.BytesIO()
    writer.write(buf)
    
    meta, chunks = service.ingest_document("scanned_receipt.pdf", buf.getvalue())
    assert meta.status == "INDEXED"
    assert meta.page_count == 1
    assert "quality_warnings" in meta.extracted_entities


def test_large_pdf_pipeline(stage3_env):
    """Test extraction and chunking of a large 10-page document preserving page numbers."""
    service: DocumentService = stage3_env["service"]
    large_pdf_bytes = _create_synthetic_pdf(text="Large Enterprise Report Section Content", num_pages=10)
    
    meta, chunks = service.ingest_document("large_enterprise_doc.pdf", large_pdf_bytes)
    assert meta.status == "INDEXED"
    assert meta.page_count == 10
    assert len(chunks) >= 10
    page_numbers = {c.page_number for c in chunks}
    assert 1 in page_numbers
    assert 10 in page_numbers


def test_docx_pipeline(stage3_env):
    """Test full DOCX ingestion with headings, tables, and metadata."""
    service: DocumentService = stage3_env["service"]
    doc = docx.Document()
    doc.add_heading("Service Level Agreement", level=1)
    doc.add_paragraph("Uptime requirement is 99.99% per calendar month.")
    
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Target"
    table.rows[1].cells[0].text = "Availability"
    table.rows[1].cells[1].text = "99.99%"
    
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()
    
    meta, chunks = service.ingest_document("sla_contract.docx", docx_bytes)
    assert meta.status == "INDEXED"
    assert meta.file_type == "docx"
    assert meta.word_count > 5
    assert len(chunks) >= 1
    combined_text = " ".join([c.text for c in chunks])
    assert "99.99%" in combined_text
    assert "| Metric | Target |" in combined_text


def test_txt_pipeline_encodings(stage3_env):
    """Test TXT ingestion with standard UTF-8 and UTF-16 encodings."""
    service: DocumentService = stage3_env["service"]
    
    # 1. UTF-8 plain text
    utf8_content = "# Project Setup\n\nRun pip install -r requirements.txt to set up the environment.".encode("utf-8")
    meta1, chunks1 = service.ingest_document("setup.txt", utf8_content)
    assert meta1.status == "INDEXED"
    assert meta1.file_type == "txt"
    assert len(chunks1) >= 1
    
    # 2. UTF-16 plain text
    utf16_content = "# Configuration\n\nDatabase connections are managed via SQLite.".encode("utf-16")
    meta2, chunks2 = service.ingest_document("config.txt", utf16_content)
    assert meta2.status == "INDEXED"
    assert "SQLite" in chunks2[0].text


def test_unsupported_file_pipeline(stage3_env):
    """Test safe rejection of unsupported file types without crashing."""
    service: DocumentService = stage3_env["service"]
    
    with pytest.raises(UnsupportedFileTypeError) as exc1:
        service.validate_file("malicious_script.exe", b"MZ\x90\x00binary")
    assert ".exe" in str(exc1.value)

    with pytest.raises(UnsupportedFileTypeError) as exc2:
        service.validate_file("run.sh", b"#!/bin/bash\necho 'hello'")
    assert ".sh" in str(exc2.value)


def test_malformed_docx_pipeline(stage3_env):
    """Test malformed DOCX file that fails zip decompression raises FileCorruptedError."""
    service: DocumentService = stage3_env["service"]
    bad_docx_bytes = b"PK\x03\x04\x14\x00\x00\x00NotARealZipArchiveAtAll\xff\xfe\x00\x01"
    
    with pytest.raises(FileCorruptedError):
        service.ingest_document("corrupted_contract.docx", bad_docx_bytes)


def test_password_protected_pdf_pipeline(stage3_env):
    """Test rejection and error tracking for password-protected encrypted PDFs."""
    service: DocumentService = stage3_env["service"]
    locked_pdf_bytes = _create_encrypted_pdf("secure_password_999")
    
    with pytest.raises(ExtractionError) as exc_info:
        service.ingest_document("confidential_payroll.pdf", locked_pdf_bytes)
    assert "password" in str(exc_info.value).lower()


def test_unusual_unicode_characters_pipeline(stage3_env):
    """Test handling and normalization of unusual Unicode characters, BOM, zero-width spaces, and emojis."""
    service: DocumentService = stage3_env["service"]
    
    raw_text = (
        "\ufeff\u200b# Multilingual Document Understanding \U0001F680\n\n"
        "Zero-width\u200b spaces are cleanly stripped.\u00a0Non-breaking space normalized.\n\n"
        "## Arabic and Cyrillic\n"
        "\u0645\u0631\u062d\u0628\u0627 \u0628\u0627\u0644\u0639\u0627\u0644\u0645 and \u041f\u0440\u0438\u0432\u0435\u0442 \u043c\u0438\u0440.\n\n"
        "## Chinese and Mathematics\n"
        "\u667a\u80fd\u6587\u6863\u5206\u6790\u7cfb\u7edf \u2211_{i=1}^n x_i \U0001F4D1"
    )
    meta, chunks = service.ingest_document("multilingual_test.txt", raw_text.encode("utf-8"))
    assert meta.status == "INDEXED"
    assert meta.word_count > 0
    assert len(chunks) >= 1
    
    full_chunk_text = " ".join([c.text for c in chunks])
    assert "\u200b" not in full_chunk_text
    assert "\ufeff" not in full_chunk_text
    assert "Multilingual Document Understanding" in full_chunk_text
    assert "\u041f\u0440\u0438\u0432\u0435\u0442" in full_chunk_text
    assert "\u667a\u80fd\u6587\u6863\u5206\u6790\u7cfb\u7edf" in full_chunk_text


def test_file_size_exceeded_rejection(stage3_env, monkeypatch):
    """Test that file exceeding maximum configured upload size is rejected safely."""
    from app.config import settings
    service: DocumentService = stage3_env["service"]
    monkeypatch.setattr(settings, "MAX_UPLOAD_SIZE_MB", 1)
    
    oversized_content = b"%PDF-1.4\n" + b"A" * (2 * 1024 * 1024)
    with pytest.raises(FileSizeExceededError):
        service.validate_file("huge_file.pdf", oversized_content)

